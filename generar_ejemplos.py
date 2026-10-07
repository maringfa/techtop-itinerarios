"""Renueva únicamente los dos XLSX ficticios incluidos en samples/."""
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook


root = Path(__file__).resolve().parent / "samples"
root.mkdir(exist_ok=True)
today = datetime.now(ZoneInfo("America/Costa_Rica")).date()

book = Workbook()
ws = book.active
ws.title = "Importaciones"
ws.append(["Items", "Status", "MBL", "HBL", "Container", "Origin Port", "Final Port",
           "ETD", "ETA", "Size", "PKG", "Navy", "DUA", "DUA DATE-HACIENDA", "ATP"])
for index in range(13):
    ws.append(["Plywood" if index % 3 == 0 else "Partes de motores", None, "", None,
               f"TTIU{1000000+index}", "", "", None, None, 40, "", "", "", "",
               today + timedelta(days=0 if index == 0 else 1)])
for start, end, number in ((3, 8, 2), (9, 14, 3)):
    ws.cell(start, 2, "In transit")
    ws.cell(start, 4, f"HBL-EJEMPLO-{number:03d}")
    ws.cell(start, 8, today - timedelta(days=6))
    ws.cell(start, 9, today + timedelta(days=2))
    for column in (2, 4, 8, 9):
        ws.merge_cells(start_row=start, start_column=column, end_row=end, end_column=column)
ws.cell(2, 2, "In transit")
ws.cell(2, 4, "HBL-EJEMPLO-001")
ws.cell(2, 8, today - timedelta(days=6))
ws.cell(2, 9, today + timedelta(days=2))
ws.cell(15, 1, "Plywood")
ws.cell(15, 2, "Arrived")
ws.cell(15, 4, "HBL-EJEMPLO-099")
ws.cell(15, 5, "TTIU1999999")
ws.cell(15, 9, today - timedelta(days=3))
ws.cell(15, 15, today)
ws.cell(16, 1, "Equipo")
ws.cell(16, 2, "In transit")
ws.cell(16, 4, "HBL-EJEMPLO-100")
ws.cell(16, 5, "TTIU1888888")
ws.cell(16, 9, today + timedelta(days=1))
for row in range(2, 17):
    for column in (8, 9, 15):
        ws.cell(row, column).number_format = "mm/dd/yyyy"
book.save(root / "importaciones_ficticias.xlsx")

book = Workbook()
ws = book.active
ws.title = "Exportaciones"
ws.append(["Dest.", "Origin", "ETD", "ETA", "Vessel", "Voyage", "Cut off Doc", "Cut off Cont",
           "Placement at the Techtop Plant", "TTI Schedule", "Deliver at TTICR", "Reservation",
           "Container", "Transfer", "ISF", "Quote"])
ws.append(["Limon", "Atlanta", "", "", "", "", "", "", "", "", today + timedelta(days=1),
           "HBL-EXP-EJEMPLO-001", "TTIU2000001", "300101", "", ""])
ws.append(["Limon", "Louisville", "", "", "", "", "", "", "", "", today + timedelta(days=1),
           "", "1 x 20", "300102 / 300103", "", ""])
ws.append(["Limon", "St. Louis", "", "", "", "", "", "", "", "", today + timedelta(days=2),
           "", "", "", "", ""])
for row in range(2, 5):
    ws.cell(row, 11).number_format = "mm/dd/yyyy"
book.save(root / "exportaciones_ficticias.xlsx")
print("Ejemplos ficticios renovados para la fecha de Costa Rica.")
