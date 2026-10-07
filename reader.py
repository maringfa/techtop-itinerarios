"""Lectura local de XLSX. Nunca escribe en los itinerarios ni registra sus valores."""
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


FIELDS = {
    "import": {"items": ("items",), "status": ("status",), "hbl": ("hbl",),
               "container": ("container",), "eta": ("eta",),
               "plant": ("atp", "llegada a planta", "llegada programada a planta", "arribo a planta")},
    "export": {"origin": ("origin",), "reservation": ("reservation",),
               "container": ("container",), "transfer": ("transfer",),
               "departure": ("deliver at tticr",)},
}


class FormatError(Exception):
    pass


def clean(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"\s+", " ", "".join(c for c in value if not unicodedata.combining(c))).strip().casefold()


def display(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_date(value, epoch, date_order="MDY"):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and value > 0:
        return from_excel(value, epoch).date() if isinstance(from_excel(value, epoch), datetime) else from_excel(value, epoch)
    value = display(value)
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y" if date_order == "MDY" else "%d/%m/%Y",
                "%m/%d/%y" if date_order == "MDY" else "%d/%m/%y"):
        try:
            return datetime.strptime(value.split(" ")[0], fmt).date()
        except ValueError:
            pass
    return None


def checked_date(value, epoch, date_order, field, row):
    parsed = as_date(value, epoch, date_order)
    if display(value) and parsed is None:
        raise FormatError(f"Fecha no reconocida en {field}, fila {row}")
    return parsed


def worksheet(book, name):
    if name:
        if name not in book:
            raise FormatError("No se encontró la hoja configurada")
        return book[name]
    return next((ws for ws in book.worksheets if ws.sheet_state == "visible"), book.active)


def headers(ws, kind, header_row=None, skip_unrelated=False):
    aliases = FIELDS[kind]
    rows = [header_row] if header_row else range(1, min(ws.max_row, 40) + 1)
    best = ({}, None)
    for number in rows:
        cells = {clean(ws.cell(number, col).value): col for col in range(1, min(ws.max_column, 100) + 1)}
        found = {field: next((cells[a] for a in names if a in cells), None) for field, names in aliases.items()}
        found = {field: col for field, col in found.items() if col}
        if len(found) > len(best[0]):
            best = (found, number)
    found, number = best
    required = set(aliases) - ({"plant"} if kind == "import" else set())
    missing = required - found.keys()
    if missing:
        core = {"container", "eta"} if kind == "import" else {"departure"}
        looks_like_itinerary = core <= found.keys() or len(found) >= 3
        if skip_unrelated and not looks_like_itinerary:
            return None, None
        raise FormatError("Faltan encabezados requeridos: " + ", ".join(sorted(missing)))
    return found, number


def merged_lookup(ws, needed_cols):
    lookup = {}
    for area in ws.merged_cells.ranges:
        if not any(area.min_col <= col <= area.max_col for col in needed_cols):
            continue
        for row in range(area.min_row, area.max_row + 1):
            for col in needed_cols:
                if area.min_col <= col <= area.max_col:
                    lookup[(row, col)] = (area.min_row, area.min_col)
    return lookup


def read_sheet(ws, epoch, kind, header_row, date_order, skip_unrelated=False):
    cols, row0 = headers(ws, kind, header_row, skip_unrelated)
    if cols is None:
        return None
    if kind == "import" and "plant" not in cols:
        raise FormatError("Falta agregar la columna ATP al itinerario de importaciones")
    merges = merged_lookup(ws, set(cols.values()))
    def val(row, field):
        col = cols[field]
        anchor = merges.get((row, col), (row, col))
        return ws.cell(*anchor).value
    result = []
    for row in range(row0 + 1, ws.max_row + 1):
        container = display(val(row, "container"))
        if kind == "import":
            if not container:
                continue
            result.append({
                "container": container, "hbl": display(val(row, "hbl")),
                "items": display(val(row, "items")), "status": display(val(row, "status")),
                "eta": checked_date(val(row, "eta"), epoch, date_order, "ETA", row),
                "plant": checked_date(val(row, "plant"), epoch, date_order, "ATP", row),
            })
        else:
            departure = checked_date(val(row, "departure"), epoch, date_order, "Deliver at TTICR", row)
            transfer = display(val(row, "transfer"))
            if not departure or not transfer:
                continue
            result.append({
                "departure": departure, "origin": display(val(row, "origin")),
                "reservation": display(val(row, "reservation")),
                "container": container, "transfer": transfer,
            })
    return result


def read_xlsx(path, kind, sheet=None, header_row=None, date_order="MDY"):
    if not Path(path).is_file():
        raise FormatError("No se encontró un archivo de itinerario configurado")
    if Path(path).suffix.lower() != ".xlsx":
        raise FormatError("Este piloto requiere archivos .xlsx")
    book = load_workbook(path, data_only=True, keep_links=False, read_only=False)
    try:
        all_sheets = sheet == "*"
        selected = book.worksheets if all_sheets else [worksheet(book, sheet)]
        result, matched = [], False
        for ws in selected:
            rows = read_sheet(ws, book.epoch, kind, header_row, date_order, skip_unrelated=all_sheets)
            if rows is not None:
                matched = True
                result.extend(rows)
        if not matched:
            raise FormatError("No se encontraron hojas con encabezados del itinerario configurado")
        return result
    finally:
        book.close()
