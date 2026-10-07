from datetime import date, datetime, time, timedelta
from pathlib import Path
import sys
import tempfile
import unittest
from zoneinfo import ZoneInfo
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from domain import build_board
from reader import FormatError, read_xlsx


ROOT = Path(__file__).resolve().parents[1]


class ItineraryTests(unittest.TestCase):
    def test_monthly_exports_include_future_months_and_new_tabs(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "monthly.xlsx"
            book = Workbook()
            book.active.title = "Resumen"
            book.active.append(["Mes", "Total"])
            for name, departures in (("Enero 2027", [date(2027, 1, 5)]),
                                     ("Octubre", [date(2026, 10, 6), date(2026, 10, 7)]),
                                     ("Septiembre", [date(2026, 9, 30)]),
                                     ("November", [date(2026, 11, 2)])):
                ws = book.create_sheet(name)
                ws.append(["Itinerario mensual"])
                ws.append(["Origin", "Reservation", "Container", "Transfer", "Deliver at TTICR"])
                for day in departures:
                    ws.append(["Atlanta", "", "", day.isoformat(), day])
            book["Enero 2027"].sheet_state = "hidden"
            book.save(path)
            now = datetime(2026, 10, 7, 12, tzinfo=ZoneInfo("America/Costa_Rica"))
            rows = read_xlsx(path, "export", "*")
            board = build_board([], rows, now=now)
            self.assertEqual([x["departure"] for x in board["exports"]],
                             ["2026-10-07", "2026-11-02", "2027-01-05"])
            self.assertEqual(len(read_xlsx(path, "export", "Octubre")), 2)
            ws = book.create_sheet("Diciembre")
            ws.append(["Origin", "Reservation", "Container", "Transfer", "Deliver at TTICR"])
            ws.append(["Louisville", "", "", "NUEVO", date(2026, 12, 2)])
            book.save(path)
            book.close()
            refreshed = build_board([], read_xlsx(path, "export", "*"), now=now)
            self.assertEqual([x["departure"] for x in refreshed["exports"]],
                             ["2026-10-07", "2026-11-02", "2026-12-02", "2027-01-05"])

    def test_incomplete_monthly_itinerary_is_not_silently_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "incomplete.xlsx"
            fields = ["Origin", "Reservation", "Container", "Transfer", "Deliver at TTICR"]
            for removed in fields:
                with self.subTest(missing=removed):
                    book = Workbook()
                    book.active.title = "Octubre"
                    book.active.append(fields)
                    ws = book.create_sheet("Noviembre")
                    ws.append([field for field in fields if field != removed])
                    book.save(path)
                    book.close()
                    with self.assertRaisesRegex(FormatError, "Faltan encabezados requeridos"):
                        read_xlsx(path, "export", "*")

    def test_wildcard_requires_at_least_one_matching_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "summary.xlsx"
            book = Workbook()
            book.active.append(["Mes", "Total"])
            book.save(path)
            book.close()
            with self.assertRaisesRegex(FormatError, "No se encontraron hojas"):
                read_xlsx(path, "export", "*")

    def test_monthly_imports_preserve_shared_cells_and_individual_dates(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "monthly-imports.xlsx"
            book = Workbook()
            october = book.active
            october.title = "Octubre"
            october.append(["Items", "Status", "HBL", "Container", "ETA", "ATP"])
            october.append(["Plywood", "In transit", "HBL-OCT", "OCT-1", date(2026, 10, 7), date(2026, 10, 8)])
            october.append(["Plywood", None, None, "OCT-2", None, None])
            for column in (2, 3, 5):
                october.merge_cells(start_row=2, end_row=3, start_column=column, end_column=column)
            november = book.create_sheet("Noviembre")
            november.append(["Itinerario"])
            november.append(["Items", "Status", "HBL", "Container", "ETA", "ATP"])
            november.append(["Equipo", "In transit", "HBL-NOV", "NOV-1", date(2026, 11, 1), None])
            book.save(path)
            book.close()
            rows = read_xlsx(path, "import", "*")
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[0]["hbl"], rows[1]["hbl"])
            self.assertIsNone(rows[1]["plant"])
            self.assertEqual(rows[2]["hbl"], "HBL-NOV")
            self.assertEqual(rows[2]["eta"], date(2026, 11, 1))

    def test_eta_is_used_only_when_atp_is_missing(self):
        now = datetime(2026, 10, 7, 12, tzinfo=ZoneInfo("America/Costa_Rica"))
        today = now.date()
        def imp(container, plant, eta, status="In transit"):
            return {"container": container, "plant": plant, "eta": eta, "status": status,
                    "hbl": "HBL-EJEMPLO", "items": "Plywood"}
        tomorrow = today + timedelta(days=1)
        imports = [imp("ATP-futuro", tomorrow, today),
                   imp("ETA-hoy", None, today),
                   imp("ATP-hoy", today, tomorrow),
                   imp("ATP-pasado", today - timedelta(days=1), tomorrow),
                   imp("ETA-pasado", None, today - timedelta(days=1)),
                   imp("Ya-arribo", None, tomorrow, "Arrived"),
                   imp("Sin-fechas", None, None)]
        board = build_board(imports, [], now=now)
        self.assertEqual([x["container"] for x in board["imports"]],
                         ["ATP-hoy", "ETA-hoy", "ATP-futuro"])
        self.assertEqual([x["displayDate"] for x in board["imports"]],
                         [today.isoformat(), today.isoformat(), tomorrow.isoformat()])
        fallback = next(x for x in board["imports"] if x["container"] == "ETA-hoy")
        self.assertEqual(fallback["dateType"], "ETA")
        self.assertIsNone(fallback["plant"])
        imports[1]["plant"] = tomorrow
        updated = next(x for x in build_board(imports, [], now=now)["imports"]
                       if x["container"] == "ETA-hoy")
        self.assertEqual(updated["dateType"], "ATP")
        self.assertEqual(updated["displayDate"], tomorrow.isoformat())
        self.assertNotIn("eta", updated)

    def test_upcoming_movements_exclude_past_and_arrived_and_sort_by_date(self):
        now = datetime(2026, 10, 7, 12, tzinfo=ZoneInfo("America/Costa_Rica"))
        today = now.date()
        def imp(days, status="In transit"):
            return {"plant": today + timedelta(days=days), "status": status,
                    "container": str(days), "hbl": "example", "items": "Plywood", "eta": today}
        def exp(days, transfer="example"):
            return {"departure": today + timedelta(days=days), "transfer": transfer,
                    "origin": "Atlanta", "container": "example", "reservation": "example"}
        board = build_board([imp(2), imp(-1), imp(1, "Arrived"), imp(0)],
                            [exp(2), exp(-1), exp(0), exp(1, "")], now=now)
        dates = [today.isoformat(), (today + timedelta(days=2)).isoformat()]
        self.assertEqual([x["plant"] for x in board["imports"]], dates)
        self.assertEqual([x["departure"] for x in board["exports"]], dates)

    def test_merged_shared_dates_and_individual_arrivals(self):
        rows = read_xlsx(ROOT / "samples/importaciones_ficticias.xlsx", "import", "Importaciones", 1)
        self.assertEqual(rows[1]["hbl"], rows[6]["hbl"])
        self.assertEqual(rows[1]["eta"], rows[6]["eta"])
        self.assertEqual(len({x["container"] for x in rows}), len(rows))
        self.assertTrue(any(x["plant"] is None for x in rows))

    def test_filters_and_overflow_data_are_not_lost(self):
        imports = read_xlsx(ROOT / "samples/importaciones_ficticias.xlsx", "import", "Importaciones", 1)
        exports = read_xlsx(ROOT / "samples/exportaciones_ficticias.xlsx", "export", "Exportaciones", 1)
        sample_today = imports[0]["plant"]
        board = build_board(imports, exports, demo=True,
                            now=datetime.combine(sample_today, time(12), ZoneInfo("America/Costa_Rica")))
        self.assertEqual(len(board["imports"]), 14)
        self.assertEqual(len(board["exports"]), 2)
        self.assertEqual(board["port"]["unscheduledCount"], 1)
        self.assertEqual(sum(x["plant"] == (sample_today + timedelta(days=1)).isoformat()
                             for x in board["imports"]), 12)
        fallback = next(x for x in board["imports"] if x["plant"] is None)
        self.assertEqual(fallback["dateType"], "ETA")
        self.assertEqual(fallback["displayDate"], (sample_today + timedelta(days=1)).isoformat())


if __name__ == "__main__":
    unittest.main()
